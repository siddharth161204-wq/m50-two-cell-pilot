#!/usr/bin/env python3
"""Pilot B results workbook for working from the numbers (one sheet per table, a read-me sheet first).

Usage: python3 make_workbook.py --out ../results   -> results/pilot_b_results.xlsx
"""
import argparse, csv, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

import lumped_two_cell as LP

ap = argparse.ArgumentParser()
ap.add_argument("--out", default="../results")
args = ap.parse_args()
O = args.out
ARIAL = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEAD_FILL = PatternFill("solid", fgColor="F0EFEC")


def num(x):
    if x in ("", None, "never"):
        return None if x != "never" else "Never"
    try:
        return float(x)
    except ValueError:
        return x[:1].upper() + x[1:] if isinstance(x, str) and x else x


def sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = BOLD; c.fill = HEAD_FILL; c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = ARIAL
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(11, min(34, len(str(h)) + 2)))
    ws.freeze_panes = "A2"
    return ws


wb = Workbook()
rd = wb.active; rd.title = "Read me"
lines = [
    "Pilot B: two LG M50 cells, resolved against lumped neighbour. Repository github.com/siddharth161204-wq/m50-two-cell-pilot",
    "Predictions commit c57ccb87e699d3dbb1230122cdba3785293627e0, public 2026-10-07 06:42:35 UTC, before any coupled resolved run.",
    "Temperatures in degrees Celsius, times in seconds, conductances in W/K. Blank or Never: the level is not reached within 900 s.",
    f"Self-heating band (Koenig, Zhao, Deng 2025, Fig. 7): {LP.T_INITIAL_BAND[0]-273.15:.1f} / {LP.T_INITIAL_MEAN-273.15:.1f} / {LP.T_INITIAL_BAND[1]-273.15:.1f} C. Internal-short band: {LP.T_ONSET_BAND[0]-273.15:.1f} / {LP.T_ONSET_MEAN-273.15:.1f} / {LP.T_ONSET_BAND[1]-273.15:.1f} C.",
    "Cell-body energies (Battery Failure Databank rev. 2, 17 fully charged LG 21700-M50 tests, open-circuit voltage 4.16 to 4.19 V): 8751.3, 21495.3, 33222.9 J (min, median, max).",
    "Sheets: Lumped (script, explicit Euler 0.01 s); Resolved two strips (42 cases); Resolved one strip (7 cases, median energy, 30 s); Scorecard; Databank rows used; Verification; Exploratory (not pre-registered: two-way coupling, isotropic roll, hot volume).",
    "Resolved max: hottest node of the resolved cell 2. Resolved mean: heat-capacity-weighted mean. Location: the path that heats the hottest node (DECISIONS.md D14).",
    "Columns in blue on 'Resolved two strips' are formulas: time saved and ratio at 183.5 C where both models cross.",
]
for i, t in enumerate(lines, 1):
    rd.cell(row=i, column=1, value=t).font = BOLD if i == 1 else ARIAL
rd.column_dimensions["A"].width = 150

lum = list(csv.DictReader(open(os.path.join(O, "lumped_results.csv"))))
hdr = ["E_body (J)", "Tau_rel (s)", "G_bus (W/K)", "G_gap (W/K)", "View factor", "T1 peak (C)", "T2 peak (C)"] + \
      [f"{k.replace('_', ' ').capitalize()} (s)" for k in LP.LEVELS]
sheet(wb, "Lumped", hdr, [[num(r[c]) for c in ["E_body_J", "tau_rel_s", "G_bus_WK", "G_gap_WK", "F_view", "T1max_C", "T2max_C"] + list(LP.LEVELS)] for r in lum])

two = list(csv.DictReader(open(os.path.join(O, "resolved_results_two.csv"))))
base = ["E_body_J", "tau_rel_s", "G_bus_WK", "T1max_C", "T2max_lumped_C", "Tmean_resolved_peak_C", "Tmax_resolved_peak_C", "t_Tmax_peak_s",
        "location_at_peak", "geometry_at_peak"]
names = ["E_body (J)", "Tau_rel (s)", "G_bus (W/K)", "Cell 1 peak (C)", "Lumped peak (C)", "Resolved mean peak (C)", "Resolved max peak (C)",
         "Time of resolved max peak (s)", "Hot spot path at peak", "Hot spot place at peak"]
cross = []
for k in LP.LEVELS:
    lab = k.replace("_", " ").capitalize()
    cross += [(f"{k}_lumped_s", f"{lab}: lumped (s)"), (f"{k}_resolved_max_s", f"{lab}: resolved max (s)"),
              (f"{k}_resolved_mean_s", f"{lab}: resolved mean (s)"), (f"{k}_location", f"{lab}: hot spot path"),
              (f"{k}_Tmax_minus_Tmean_K", f"{lab}: max minus mean at crossing (K)")]
hdr = names + [c[1] for c in cross] + ["Time saved at 183.5 C (s)", "Lumped over resolved time at 183.5 C", "Energy residual"]
rows = []
for r in two:
    rows.append([num(r[c]) for c in base] + [num(r[c[0]]) for c in cross] + [None, None, num(r["energy_residual_rel"])])
ws = sheet(wb, "Resolved two strips", hdr, rows)
col = {h: i + 1 for i, h in enumerate(hdr)}
L, Rm = get_column_letter(col["T initial mean: lumped (s)"]), get_column_letter(col["T initial mean: resolved max (s)"])
for i in range(2, len(rows) + 2):
    a = ws.cell(row=i, column=col["Time saved at 183.5 C (s)"], value=f'=IF(AND(ISNUMBER({L}{i}),ISNUMBER({Rm}{i})),{L}{i}-{Rm}{i},"")')
    b = ws.cell(row=i, column=col["Lumped over resolved time at 183.5 C"], value=f'=IF(AND(ISNUMBER({L}{i}),ISNUMBER({Rm}{i})),{L}{i}/{Rm}{i},"")')
    for c in (a, b):
        c.font = Font(name="Arial", size=10, color="0000FF"); c.number_format = "0.00"

one = list(csv.DictReader(open(os.path.join(O, "resolved_results_one.csv"))))
sheet(wb, "Resolved one strip", names + [c[1] for c in cross], [[num(r[c]) for c in base] + [num(r[c[0]]) for c in cross] for r in one])

sc = json.load(open(os.path.join(O, "prediction_scorecard.json")))
pred = {"P1": "At 5 mW/K neither model reaches 169.0 C in any of the six energy and duration cases",
        "P2": "At 500 mW/K, where the lumped node crosses 183.5 C, the resolved hottest point crosses it more than 50 % earlier",
        "P3": "Only-resolved crossings exist; smallest conductance for 183.5 C is 0.1 or 0.2 W/K; at 0.5 W/K and 33.2 kJ the hottest point reaches 255.7 C",
        "P4": "Hottest point (at its peak, median energy, 30 s) gap-facing at 0.005 and 0.01 W/K, terminal end at 0.1, 0.2 and 0.5 W/K",
        "P5": "At 0.5 W/K the resolved mean's peak rise is at least 15 % smaller than the lumped node's, in all six cases"}
meas = {"P1": "Peaks 39.3 to 77.0 C (lumped and resolved)",
        "P2": "; ".join(f"Tau {int(c['tau'])} s: {c['resolved_max_s']:.2f} s against {c['lumped_s']:.2f} s, {100*c['earlier_fraction']:.1f} % earlier" for c in sc["P2"]["cases"]),
        "P3": f"Smallest conductance {sc['P3']['smallest_G_only_resolved_T_initial_mean']} W/K; 255.7 C reached at " +
              " and ".join(f"{float(r['T_onset_low_resolved_max_s']):.2f} s (tau {int(float(r['tau_rel_s']))} s)" for r in two
                           if float(r['G_bus_WK']) == 0.5 and abs(float(r['E_body_J']) - 33222.9) < 1 and r['T_onset_low_resolved_max_s']),
        "P4": "; ".join(f"{g}: {v}" for g, v in sc["P4"]["labels_at_peak"].items()),
        "P5": "; ".join(f"{c['E']/1e3:.1f} kJ, {int(c['tau'])} s: {100*c['reduction']:.1f} %" for c in sc["P5"]["cases"])}
sheet(wb, "Scorecard", ["Prediction", "Statement (predictions.md)", "Result", "Measured"],
      [[k, pred[k], "Held" if sc[k]["held"] else "Failed", meas[k]] for k in ("P1", "P2", "P3", "P4", "P5")],
      widths={1: 11, 2: 70, 3: 9, 4: 90})

db = list(csv.DictReader(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "inputs", "databank", "databank_m50_used.csv"))))
sheet(wb, "Databank rows used", list(db[0].keys()), [[num(v) for v in r.values()] for r in db])

N = json.load(open(os.path.join(O, "note_numbers.json")))
ver = [["V1 radiation only: max deviation of resolved mean from lumped node (% of rise)", N["V1"]["max_rel_dev_percent"]],
       ["V1 energy residual (relative)", N["V1"]["energy_residual"]],
       ["V2 median case (21.5 kJ, 30 s, 50 mW/K), time step halved: change in resolved max peak (K)", N["V2_median_dt_half_peak_changes_K"]["dt_half_Tmax"]],
       ["V2 median case, time step halved: change in resolved mean peak (K)", N["V2_median_dt_half_peak_changes_K"]["dt_half_Tmean"]]]
if "V2_median_refined" in N:
    ver += [["V2 median case, mesh refined (every spacing halved): change in resolved max peak (K)", N["V2_median_refined"]["peak_Tmax_change_K"]],
            ["V2 median case, mesh refined: change in resolved mean peak (K)", N["V2_median_refined"]["peak_Tmean_change_K"]]]
for key in sorted(k for k in N if k.startswith("v2_")):
    v = N[key]
    _, en, tau, g, ref, dt = key.split("_")
    what = f"mesh refined ({v['nodes']:,} nodes)" if ref != "refine1" else f"time step halved ({v['dt_s']:g} s)"
    lab = f"V2 {en} energy, {tau.replace('tau', 'tau ')} s, {float(g[1:]) * 1000:g} mW/K, {what}"
    lab += " (against an unrounded base run)" if v.get("reference", "").startswith("unrounded") else " (against the sweep table, rounded to 0.01 s)"
    ver += [[f"{lab}: largest change in a hottest-point crossing time (s)", N["V2_max_abs_change_hottest_crossing_s"][key]],
            [f"{lab}: change in resolved max peak (K)", v["peak_Tmax_change_K"]],
            [f"{lab}: change in resolved mean peak (K)", v["peak_Tmean_change_K"]]]
for key, v in N.get("V2_base_runs", {}).items():
    ver.append([f"Unrounded rerun of the base case ({key}): crossing times agree with the sweep table to its 0.01 s rounding",
                "Yes" if v["matches_sweep_table_to_0.01_s"] else "No"])
for key, v in N.get("rerun_check", {}).items():
    ver += [[f"Rerun of {key} with the final code: fields differing from the registered row, run time excepted", len(v["fields_differing_except_runtime"])],
            [f"Rerun of {key} with the final code: largest difference in any recorded time series", v["max_abs_series_difference"]]]
ver += [["V3 median case: lumped peak (C)", N["V3"]["peak_lumped_C"]], ["V3 median case: resolved mean peak (C)", N["V3"]["peak_resolved_mean_C"]],
        ["V3 median case: resolved max peak (C)", N["V3"]["peak_resolved_max_C"]],
        ["Largest energy residual over the 42 cases (relative)", N["energy_residual_max_abs"]]]
cw = os.path.join(O, "check_two_way.json")
if os.path.exists(cw):
    c = json.load(open(cw))
    ver += [["Two-way coupling check (roll at 1e4 W/mK, 33.2 kJ, 30 s, 500 mW/K): largest difference in cell 1 from the lumped pair (K)", c["max_abs_T1_diff_K"]],
            ["Two-way coupling check: largest difference in cell 2 mean from the lumped pair (K)", c["max_abs_T2_diff_K"]]]
wsv = sheet(wb, "Verification", ["Check", "Value"], ver, widths={1: 110, 2: 16})
for (c,) in wsv.iter_rows(min_row=2, min_col=2, max_col=2):
    if isinstance(c.value, float) and c.value != 0 and abs(c.value) < 1e-4:
        c.number_format = "0.0E+00"

# exploratory runs, not pre-registered (DECISIONS.md D23, D25, D26)
ex_rows = []
for mode, lab in (("sensitivity_two_way", "Two-way coupling"), ("sensitivity_isotropic", "Isotropic roll, 25 W/mK")):
    for key, v in N.get(mode, {}).items():
        en, tau, g = key.split("_")
        ex_rows.append([lab, {"max": 33222.9, "median": 21495.3, "min": 8751.3}[en], float(tau[3:]), float(g[1:]),
                        v["t183_hottest_one_way_s"], v["t183_hottest_s"], v["peak_hottest_one_way_C"], v["peak_hottest_C"],
                        v["peak_mean_one_way_C"], v["peak_mean_C"], v["t169_mean_s"] if v["t169_mean_s"] is not None else "Never",
                        v["lumped_peak_C"], v["mean_rise_reduction_percent"], v["T1_peak_lumped_pair_C"], v["T1_peak_C"],
                        num(v["location_at_peak"])])
ws = sheet(wb, "Exploratory", ["Run", "E_body (J)", "Tau_rel (s)", "G_bus (W/K)", "183.5 C, hottest point, registered (s)",
                               "183.5 C, hottest point, this run (s)", "Hottest peak, registered (C)", "Hottest peak, this run (C)",
                               "Mean peak, registered (C)", "Mean peak, this run (C)", "169.0 C, mean, this run (s)", "Lumped node peak (C)",
                               "Mean rise below lumped rise, this run (%)", "Cell 1 peak, lumped pair (C)", "Cell 1 peak, this run (C)",
                               "Hot spot path at peak"], ex_rows, widths={1: 24})
r0 = len(ex_rows) + 4
hv = N.get("hot_volume", {})
hdr = ["Hot volume run (registered physics)", "E_body (J)", "Tau_rel (s)", "G_bus (W/K)", "Max % of cell at or above 169.0 C",
       "Max % at or above 183.5 C", "Max % at or above 198.0 C", "Max % at or above 255.7 C", "0.1 % of cell at 183.5 C (s)",
       "1 % of cell at 183.5 C (s)"]
for j, h in enumerate(hdr, 1):
    c = ws.cell(row=r0, column=j, value=h); c.font = BOLD; c.fill = HEAD_FILL; c.alignment = Alignment(wrap_text=True, vertical="top")
for i, (key, v) in enumerate(hv.items(), 1):
    _, en, tau, g = key.split("_")
    vals = [f"Two strips, {en} energy, tau {tau[3:]} s, {float(g[1:]) * 1000:g} mW/K", v["E_body_J"], v["tau_rel_s"], v["G_bus_WK"],
            v["max_volume_percent_above_169.0"], v["max_volume_percent_above_183.5"], v["max_volume_percent_above_198.0"],
            v["max_volume_percent_above_255.7"], v.get("t_0.1pct_above_183.5_s") or "Never", v.get("t_1pct_above_183.5_s") or "Never"]
    for j, x in enumerate(vals, 1):
        c = ws.cell(row=r0 + i, column=j, value=x); c.font = ARIAL
        if isinstance(x, float) and j >= 5:
            c.number_format = "0.000"
wb.save(os.path.join(O, "pilot_b_results.xlsx"))
print("wrote", os.path.join(O, "pilot_b_results.xlsx"))
